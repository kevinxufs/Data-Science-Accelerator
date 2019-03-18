#Follows the exact same process as air_processing


import pandas as pd
import os
import openpyxl
from openpyxl import load_workbook
from openpyxl import Workbook
from pandas import ExcelWriter
from pandas import ExcelFile

datasets = ['hs2_noise_data_camden_february_2018.xlsx',  'hs2_noise_data_october_2017.xlsx',
 'hs2_noise_data_ealing_july_2018.xlsx',
 'hs2_noise_data_camden_july_2018.xlsx',
 'hs2_noise_data_camden_april_2018.xlsx',
 'hs2_noise_data_ealing_may_2018.xlsx',
 'hs2_noise_data_ealing_march_2018.xlsx',
 'hs2_noise_data_camden_may_2018.xlsx',
 'hs2_noise_data_ealing_april_2018.xlsx',
 'hs2_noise_data_camden_june_2018.xlsx',
 'hs2_noise_data_camden_jan_2018.xlsx',
 'hs2_noise_data_ealing_june_2018.xlsx',
 'hs2_noise_data-december_2017.xlsx',
 'hs2_noise_data_november_2017.xlsx',
 'hs2_noise_data_september_2017.xlsx',
 'hs2_noise_data_camden_march_2018.xlsx']


os.chdir('/Users/datascience9/Desktop/noise quality')
metarefdata = load_workbook('metarefdatanoise.xlsx')

x = metarefdata.active
all_rows = []

for row in x:
    current_row = []
    for cell in row:
        current_row.append(cell.value)
    all_rows.append(current_row)
    
completedic = {}

for lists in all_rows:
    completedic[lists[0]] = lists[1:]
    
col = ['Date/Time', "Period, T", 'LpAeq,T', "LpAF,Max", "LpA90,T", 'Monitor Ref', 'Location 1', 'Location 2', 'GPS reference 1', 'GPS reference 2', 'Grid reference', 'Latitude', 'Longitude', 'Postcode', 'Worksite Ref']

completedf = pd.DataFrame(columns = col)

for file in datasets:
    wb = load_workbook(file)
    sheets = wb.sheetnames[1:len(wb.sheetnames)-1]
    print('Now processing file ' + file)
    for i in sheets:
        df = pd.read_excel(file, sheet_name = i, skiprows = 3)
        ref = wb[i]['B1'].value
        df['Monitor Ref'] = ref
        df['Location 1'] = completedic[ref][0]
        df['Location 2'] = completedic[ref][1]
        df['GPS reference 1'] = completedic[ref][2]
        df['GPS reference 2'] = completedic[ref][3]
        df['Grid reference'] = completedic[ref][4]
        df['Latitude'] = completedic[ref][5]
        df['Longitude'] = completedic[ref][6]
        df['Postcode'] = completedic[ref][7]
        df['Worksite Ref'] = completedic[ref][8]
        completedf = pd.concat([completedf, df])
writer = pd.ExcelWriter('completenoisedatav2.xlsx', engine ='xlsxwriter')

completedf.to_excel(writer, sheet_name = 'Sheet1', index = False)

writer.save()
